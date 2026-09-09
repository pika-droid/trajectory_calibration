"""Export zero-shot cross-domain transfer (LODO) and intercept adaptation tables to an Excel file.

Sheet 1: platt_5d_cross_domain (NC, Temp, Platt 1D, Platt 5D across All 14 + 8 Winning Datasets)
Sheet 2: vcps_5d_cross_domain (NC, Temp, Platt 1D, VCPS 5D across All 14 + 6 Winning Datasets)
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

ALL_14_DATASETS: list[str] = [
    "ai2d",
    "chartqa",
    "docvqa",
    "gqa",
    "infographicvqa",
    "lego-puzzles",
    "mmbench",
    "mmmu",
    "pope",
    "scienceqa",
    "seedbench",
    "textvqa",
    "vizwiz-vqa",
    "vqav2",
]

PLATT_5D_DATASETS: list[str] = [
    "chartqa",
    "infographicvqa",
    "lego-puzzles",
    "mmbench",
    "scienceqa",
    "seedbench",
    "textvqa",
    "vizwiz-vqa",
]

VCPS_5D_DATASETS: list[str] = [
    "lego-puzzles",
    "mmbench",
    "scienceqa",
    "textvqa",
    "vizwiz-vqa",
    "vqav2",
]

# (Display Method, Raw Method, Raw Protocol, Display Protocol)
PLATT_5D_ROWS: list[tuple[str, str, str, str]] = [
    ("NC", "Naive Confidence (NC)", "Zero-Shot Base", "Zero-Shot Base"),
    ("Temp", "Temperature Scaling (TS)", "Zero-Shot Base", "Zero-Shot Base"),
    ("Platt 1D", "Platt Scaling (1D)", "Zero-Shot Base", "Zero-Shot Base"),
    ("Platt 1D", "Platt Scaling (1D)", "Target Adapted (Saerens-EM)", "Intercept Adapted"),
    ("Platt 5D", "Trajectory Platt (5D)", "Zero-Shot Base", "Zero-Shot Base"),
    ("Platt 5D", "Trajectory Platt (5D)", "Target Adapted (Saerens-EM)", "Intercept Adapted"),
]

VCPS_5D_ROWS: list[tuple[str, str, str, str]] = [
    ("NC", "Naive Confidence (NC)", "Zero-Shot Base", "Zero-Shot Base"),
    ("Temp", "Temperature Scaling (TS)", "Zero-Shot Base", "Zero-Shot Base"),
    ("Platt 1D", "Platt Scaling (1D)", "Zero-Shot Base", "Zero-Shot Base"),
    ("Platt 1D", "Platt Scaling (1D)", "Target Adapted (Saerens-EM)", "Intercept Adapted"),
    ("VPCS Platt 5D", "VCPS-5D (Our Method)", "Zero-Shot Base", "Zero-Shot Base"),
    ("VPCS Platt 5D", "VCPS-5D (Our Method)", "Target Adapted (Saerens-EM)", "Intercept Adapted"),
]


def get_cell(ws: Worksheet, row: int, col: int) -> Cell:
    """Safely retrieves a strongly-typed Cell instance from a worksheet."""
    cell = ws.cell(row=row, column=col)
    assert isinstance(cell, Cell)
    return cell


def write_cross_domain_block(
    ws: Worksheet,
    start_row: int,
    section_title: str,
    target_datasets: list[str],
    row_specs: list[tuple[str, str, str, str]],
    m3_df: pd.DataFrame,
    mqt_df: pd.DataFrame,
) -> int:
    """Writes a 7-column cross-domain block (Model, Method, Protocol, ECE, Ada-ECE, AUROC, Brier)."""
    get_cell(ws, start_row, 1).value = section_title

    headers = ["Model", "Method", "Transfer Protocol", "ECE (%)", "Ada-ECE (%)", "AUROC", "Brier"]
    for col_idx, h in enumerate(headers, start=1):
        get_cell(ws, start_row + 2, col_idx).value = h

    cur_row = start_row + 3
    # M3-LLaVA
    for disp_m, raw_m, raw_p, disp_p in row_specs:
        get_cell(ws, cur_row, 1).value = "M3-LLaVA"
        get_cell(ws, cur_row, 2).value = disp_m
        get_cell(ws, cur_row, 3).value = disp_p

        sub: Any = m3_df[
            (m3_df["test_dataset"].isin(target_datasets))
            & (m3_df["method"] == raw_m)
            & (m3_df["transfer_mode"] == raw_p)
        ]
        c_ece = get_cell(ws, cur_row, 4)
        c_ece.value = float(sub["ece"].mean())
        c_ece.number_format = "0.00%"

        c_aece = get_cell(ws, cur_row, 5)
        c_aece.value = float(sub["adaptive_ece"].mean())
        c_aece.number_format = "0.00%"

        c_auc = get_cell(ws, cur_row, 6)
        c_auc.value = float(sub["auroc"].mean())
        c_auc.number_format = "0.000"

        c_brier = get_cell(ws, cur_row, 7)
        c_brier.value = float(sub["brier"].mean())
        c_brier.number_format = "0.0000"
        cur_row += 1

    cur_row += 1  # Blank row separator

    # MQT-LLaVA
    for disp_m, raw_m, raw_p, disp_p in row_specs:
        get_cell(ws, cur_row, 1).value = "MQT-LLaVA"
        get_cell(ws, cur_row, 2).value = disp_m
        get_cell(ws, cur_row, 3).value = disp_p

        sub_mqt: Any = mqt_df[
            (mqt_df["test_dataset"].isin(target_datasets))
            & (mqt_df["method"] == raw_m)
            & (mqt_df["transfer_mode"] == raw_p)
        ]
        c_ece = get_cell(ws, cur_row, 4)
        c_ece.value = float(sub_mqt["ece"].mean())
        c_ece.number_format = "0.00%"

        c_aece = get_cell(ws, cur_row, 5)
        c_aece.value = float(sub_mqt["adaptive_ece"].mean())
        c_aece.number_format = "0.00%"

        c_auc = get_cell(ws, cur_row, 6)
        c_auc.value = float(sub_mqt["auroc"].mean())
        c_auc.number_format = "0.000"

        c_brier = get_cell(ws, cur_row, 7)
        c_brier.value = float(sub_mqt["brier"].mean())
        c_brier.number_format = "0.0000"
        cur_row += 1

    return cur_row


def write_cross_domain_sheet(
    ws: Worksheet,
    subset_label: str,
    target_subset: list[str],
    row_specs: list[tuple[str, str, str, str]],
    m3_lodo: pd.DataFrame,
    mqt_lodo: pd.DataFrame,
) -> None:
    """Writes Part A (All 14 Datasets) and Part B (Winning Subset) for cross-domain transfer."""
    title_a = "Part A: Zero-Shot Cross-Domain Transfer (LODO) - Evaluated on All 14 Datasets (Macro-Average)"
    next_row = write_cross_domain_block(
        ws, 1, title_a, ALL_14_DATASETS, row_specs, m3_lodo, mqt_lodo
    )

    title_b = (
        f"Part B: Zero-Shot Cross-Domain Transfer (LODO) - Evaluated on {subset_label} "
        f"({len(target_subset)} Datasets): {', '.join(target_subset)}"
    )
    write_cross_domain_block(ws, next_row + 2, title_b, target_subset, row_specs, m3_lodo, mqt_lodo)


def generate_cross_domain_workbook(repo_root: Path, output_path: Path) -> None:
    """Generates the 2-sheet cross-domain Excel workbook with NC, Temp, Platt, and VCPS."""
    m3_lodo = pd.read_csv(repo_root / "results/experiments/lodo/lodo_m3_summary.csv")
    mqt_lodo = pd.read_csv(repo_root / "results/experiments/lodo/lodo_mqt_summary.csv")

    wb = openpyxl.Workbook()

    # Sheet 1: Platt 5D Cross-Domain Transfer (LODO)
    ws1: Any = wb.active
    assert ws1 is not None
    ws1.title = "platt_5d_cross_domain"
    write_cross_domain_sheet(
        ws1, "Winning 8-Dataset Subset", PLATT_5D_DATASETS, PLATT_5D_ROWS, m3_lodo, mqt_lodo
    )

    # Sheet 2: VCPS 5D Cross-Domain Transfer (LODO)
    ws2 = wb.create_sheet(title="vcps_5d_cross_domain")
    write_cross_domain_sheet(
        ws2, "Winning 6-Dataset Subset", VCPS_5D_DATASETS, VCPS_5D_ROWS, m3_lodo, mqt_lodo
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"Generated cross-domain workbook with NC and TS: {output_path}")


def main() -> None:
    repo_root = Path(__file__).resolve().parent.parent
    out_file = repo_root / "results/cross_domain_tables.xlsx"
    generate_cross_domain_workbook(repo_root, out_file)


if __name__ == "__main__":
    main()
