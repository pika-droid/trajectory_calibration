"""Export trajectory calibration results to plain, unstyled Excel workbooks.

Creates simple tables without colors, styling, or banners:
- Single-model simple workbooks (calibration_results_simple_m3.xlsx, calibration_results_simple_mqt.xlsx)
- Master simple workbook with both M3-LLaVA and MQT-LLaVA (calibration_results_simple_master.xlsx)
- Dedicated Platt (5D) top-performing subset workbook (calibration_results_mean_platt5d.xlsx)
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

TARGET_METHODS: list[tuple[str, str, str]] = [
    ("NC", "Naive Confidence (NC)", "bench"),
    ("TS", "Temperature Scaling (TS)", "bench"),
    ("Platt 1d", "Platt Scaling (1D)", "bench"),
    ("Platt (5d)", "Trajectory Platt (5D)", "bench"),
    ("Platt (17D)", "Trajectory Platt (17D)", "bench"),
    ("VCPS (5D)", "VCPS-5D (Our Method)", "bench"),
    ("VCPS (17D)", "VCPS-17D (Our Method)", "bench"),
    ("Umpire", "umpire", "umpire"),
    ("Eigen score", "eigen_score", "umpire"),
    ("LN entropy", "ln_entropy", "umpire"),
    ("Semantic entropy", "semantic_entropy", "umpire"),
]

PLATT5D_DATASETS: list[str] = [
    "chartqa",
    "infographicvqa",
    "lego-puzzles",
    "mmbench",
    "scienceqa",
    "seedbench",
    "textvqa",
    "vizwiz-vqa",
]


def get_cell(ws: Worksheet, row: int, col: int) -> Cell:
    """Safely retrieves a strongly-typed Cell instance from a worksheet."""
    cell = ws.cell(row=row, column=col)
    assert isinstance(cell, Cell)
    return cell


def write_single_model_block(
    ws: Worksheet,
    start_row: int,
    model_tag: str,
    dataset_name: str,
    bench_df: pd.DataFrame,
    ump_df: pd.DataFrame,
    include_model_col: bool = True,
) -> int:
    """Writes an unformatted block of 11 methods for a single model and dataset."""
    b_sub: Any = bench_df[bench_df["dataset"] == dataset_name]
    u_sub: Any = ump_df[ump_df["dataset"] == dataset_name]

    col_offset = 1 if include_model_col else 0

    for idx, (disp_m, raw_m, src) in enumerate(TARGET_METHODS):
        cur_row = start_row + idx
        if include_model_col:
            get_cell(ws, cur_row, 1).value = model_tag

        get_cell(ws, cur_row, 1 + col_offset).value = disp_m
        c_ece = get_cell(ws, cur_row, 2 + col_offset)
        c_aece = get_cell(ws, cur_row, 3 + col_offset)
        c_auc = get_cell(ws, cur_row, 4 + col_offset)
        c_brier = get_cell(ws, cur_row, 5 + col_offset)

        if src == "bench":
            row_data: Any = b_sub[b_sub["method"] == raw_m]
            if len(row_data) > 0:
                r = row_data.iloc[0]
                c_ece.value = float(r["ece"])
                c_ece.number_format = "0.00%"
                c_aece.value = float(r["adaptive_ece"])
                c_aece.number_format = "0.00%"
                c_auc.value = float(r["auroc"])
                c_auc.number_format = "0.000"
                c_brier.value = float(r["brier"])
                c_brier.number_format = "0.0000"
        else:
            row_data = u_sub[u_sub["method"] == raw_m]
            if len(row_data) > 0:
                r = row_data.iloc[0]
                c_ece.value = float(r["cece"])
                c_ece.number_format = "0.00%"
                c_aece.value = "-"
                c_auc.value = float(r["auc"])
                c_auc.number_format = "0.000"
                c_brier.value = "-"

    return start_row + len(TARGET_METHODS)


def write_macro_model_block(
    ws: Worksheet,
    start_row: int,
    model_tag: str,
    datasets: list[str],
    include_model_col: bool = True,
) -> int:
    """Writes an unformatted block of 11 macro formulas for a single model."""
    col_offset = 1 if include_model_col else 0

    for idx, (disp_m, _raw_m, src) in enumerate(TARGET_METHODS):
        cur_row = start_row + idx
        if include_model_col:
            get_cell(ws, cur_row, 1).value = model_tag

        get_cell(ws, cur_row, 1 + col_offset).value = disp_m
        c_ece = get_cell(ws, cur_row, 2 + col_offset)
        c_aece = get_cell(ws, cur_row, 3 + col_offset)
        c_auc = get_cell(ws, cur_row, 4 + col_offset)
        c_brier = get_cell(ws, cur_row, 5 + col_offset)

        col_c = "C" if include_model_col else "B"
        col_d = "D" if include_model_col else "C"
        col_e = "E" if include_model_col else "D"
        col_f = "F" if include_model_col else "E"

        ece_refs = ", ".join([f"'{d}'!{col_c}{cur_row}" for d in datasets])
        c_ece.value = f"=AVERAGE({ece_refs})"
        c_ece.number_format = "0.00%"

        auc_refs = ", ".join([f"'{d}'!{col_e}{cur_row}" for d in datasets])
        c_auc.value = f"=AVERAGE({auc_refs})"
        c_auc.number_format = "0.000"

        if src == "bench":
            aece_refs = ", ".join([f"'{d}'!{col_d}{cur_row}" for d in datasets])
            c_aece.value = f"=AVERAGE({aece_refs})"
            c_aece.number_format = "0.00%"

            brier_refs = ", ".join([f"'{d}'!{col_f}{cur_row}" for d in datasets])
            c_brier.value = f"=AVERAGE({brier_refs})"
            c_brier.number_format = "0.0000"
        else:
            c_aece.value = "-"
            c_brier.value = "-"

    return start_row + len(TARGET_METHODS)


def generate_simple_single_workbook(
    arch: str,
    repo_root: Path,
    output_path: Path,
) -> None:
    """Generates a plain, 5-column Excel file for a single architecture."""
    bench_df = pd.read_csv(
        repo_root / f"results/experiments/benchmark/benchmark_{arch}_summary.csv"
    )
    ump_df = pd.read_csv(repo_root / f"results/umpire_eval/{arch}_llava_cumulative_summary.csv")
    datasets = sorted(bench_df["dataset"].unique())

    wb = openpyxl.Workbook()
    ws_macro: Any = wb.active
    assert ws_macro is not None
    ws_macro.title = "Macro_Average"

    headers = ["Method", "ECE (%)", "Ada-ECE (%)", "AUROC", "Brier"]
    for col_idx, h in enumerate(headers, start=1):
        get_cell(ws_macro, 1, col_idx).value = h
    write_macro_model_block(ws_macro, 2, arch, datasets, include_model_col=False)

    for d in datasets:
        ws_d = wb.create_sheet(title=d)
        for col_idx, h in enumerate(headers, start=1):
            get_cell(ws_d, 1, col_idx).value = h
        write_single_model_block(ws_d, 2, arch, d, bench_df, ump_df, include_model_col=False)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"Generated simple single workbook: {output_path}")


def generate_simple_subset_workbook(
    repo_root: Path,
    datasets: list[str],
    summary_sheet_name: str,
    output_path: Path,
) -> None:
    """Generates a plain Excel file for a specific subset of datasets with both M3 and MQT."""
    m3_bench = pd.read_csv(repo_root / "results/experiments/benchmark/benchmark_m3_summary.csv")
    m3_ump = pd.read_csv(repo_root / "results/umpire_eval/m3_llava_cumulative_summary.csv")
    mqt_bench = pd.read_csv(repo_root / "results/experiments/benchmark/benchmark_mqt_summary.csv")
    mqt_ump = pd.read_csv(repo_root / "results/umpire_eval/mqt_llava_cumulative_summary.csv")

    wb = openpyxl.Workbook()
    ws_macro: Any = wb.active
    assert ws_macro is not None
    ws_macro.title = summary_sheet_name

    headers = ["Model", "Method", "ECE (%)", "Ada-ECE (%)", "AUROC", "Brier"]
    for col_idx, h in enumerate(headers, start=1):
        get_cell(ws_macro, 1, col_idx).value = h

    # M3-LLaVA: rows 2 to 12
    end_m3 = write_macro_model_block(ws_macro, 2, "M3-LLaVA", datasets, include_model_col=True)
    # MQT-LLaVA: rows 14 to 24 (row 13 is empty separator)
    _ = write_macro_model_block(ws_macro, end_m3 + 1, "MQT-LLaVA", datasets, include_model_col=True)

    for d in datasets:
        ws_d = wb.create_sheet(title=d)
        for col_idx, h in enumerate(headers, start=1):
            get_cell(ws_d, 1, col_idx).value = h
        end_d_m3 = write_single_model_block(
            ws_d, 2, "M3-LLaVA", d, m3_bench, m3_ump, include_model_col=True
        )
        _ = write_single_model_block(
            ws_d, end_d_m3 + 1, "MQT-LLaVA", d, mqt_bench, mqt_ump, include_model_col=True
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"Generated subset workbook: {output_path}")


def main() -> None:
    repo_root = Path(__file__).resolve().parent.parent
    all_datasets = sorted(
        pd.read_csv(repo_root / "results/experiments/benchmark/benchmark_m3_summary.csv")[
            "dataset"
        ].unique()
    )

    # 1. Master simple workbook (all 14 datasets, both M3 and MQT)
    generate_simple_subset_workbook(
        repo_root,
        all_datasets,
        summary_sheet_name="Macro_Average",
        output_path=repo_root / "results/calibration_results_simple_master.xlsx",
    )

    # 2. Platt (5D) top-performing subset workbook (8 datasets, both M3 and MQT)
    generate_simple_subset_workbook(
        repo_root,
        PLATT5D_DATASETS,
        summary_sheet_name="mean_platt5d",
        output_path=repo_root / "results/calibration_results_mean_platt5d.xlsx",
    )

    # 3. Individual simple workbooks across all 14 datasets
    generate_simple_single_workbook(
        "m3", repo_root, repo_root / "results/calibration_results_simple_m3.xlsx"
    )
    generate_simple_single_workbook(
        "mqt", repo_root, repo_root / "results/calibration_results_simple_mqt.xlsx"
    )


if __name__ == "__main__":
    main()
