"""Export trajectory calibration and baseline results to professional Excel workbooks.

Creates structured multi-tab spreadsheets for M3-LLaVA and MQT-LLaVA:
- 'Macro_Average': Overall macro mean across all 14 datasets for 4 metrics
- 14 separate dataset sheets ('ai2d', 'chartqa', ...): Individual benchmark results
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

TARGET_METHODS: list[tuple[str, str, str, str]] = [
    ("Naive Confidence (NC)", "Naive Confidence (NC)", "bench", "Single-Pass ($T=0.0, K=1$)"),
    ("Temperature Scaling (TS)", "Temperature Scaling (TS)", "bench", "Single-Pass ($T=0.0, K=1$)"),
    ("Platt Scaling (1D)", "Platt Scaling (1D)", "bench", "Single-Pass ($T=0.0, K=1$)"),
    ("Trajectory Platt (5D)", "Trajectory Platt (5D)", "bench", "Single-Pass ($T=0.0, K=1$)"),
    ("Trajectory Platt (17D)", "Trajectory Platt (17D)", "bench", "Single-Pass ($T=0.0, K=1$)"),
    ("VCPS-5D (Our Method)", "VCPS-5D (Our Method)", "bench", "Single-Pass ($T=0.0, K=1$)"),
    ("VCPS-17D (Our Method)", "VCPS-17D (Our Method)", "bench", "Single-Pass ($T=0.0, K=1$)"),
    ("Umpire", "umpire", "umpire", "Multi-Pass ($T=0.5, K=10$)"),
    ("Eigen Score", "eigen_score", "umpire", "Multi-Pass ($T=0.5, K=10$)"),
    ("LN Entropy", "ln_entropy", "umpire", "Multi-Pass ($T=0.5, K=10$)"),
    ("Semantic Entropy", "semantic_entropy", "umpire", "Multi-Pass ($T=0.5, K=10$)"),
]

FONT_BANNER = Font(name="Arial", size=11, bold=True, color="FFFFFF")
FONT_HEADER = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FONT_DATA = Font(name="Arial", size=10, bold=False, color="000000")

FILL_NAVY = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
FILL_HEADER = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
FILL_ZEBRA = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
FILL_WHITE = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

BORDER_THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
BORDER_HEADER = Border(
    left=Side(style="thin", color="FFFFFF"),
    right=Side(style="thin", color="FFFFFF"),
    top=Side(style="medium", color="1F4E79"),
    bottom=Side(style="medium", color="1F4E79"),
)


def get_cell(ws: Worksheet, row: int, col: int) -> Cell:
    """Safely retrieves a strongly-typed Cell instance from a worksheet."""
    cell = ws.cell(row=row, column=col)
    assert isinstance(cell, Cell)
    return cell


def load_all_data(repo_root: Path) -> dict[str, dict[str, pd.DataFrame]]:
    """Loads benchmark and umpire cumulative summaries for both architectures."""
    data: dict[str, dict[str, pd.DataFrame]] = {}
    for arch in ["m3", "mqt"]:
        bench_csv = repo_root / f"results/experiments/benchmark/benchmark_{arch}_summary.csv"
        ump_csv = repo_root / f"results/umpire_eval/{arch}_llava_cumulative_summary.csv"
        data[arch] = {
            "bench": pd.read_csv(bench_csv),
            "umpire": pd.read_csv(ump_csv),
        }
    return data


def write_table_headers(ws: Worksheet, start_row: int, model_name: str, title_suffix: str) -> int:
    """Writes the table banner and column headers."""
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=6)
    banner_cell = get_cell(ws, start_row, 1)
    banner_cell.value = f"{model_name} | {title_suffix}"
    banner_cell.font = FONT_BANNER
    banner_cell.fill = FILL_NAVY
    banner_cell.alignment = ALIGN_LEFT
    ws.row_dimensions[start_row].height = 24

    headers = [
        "Calibration Method",
        "Regime / Sampling",
        "ECE (%) \u2193",
        "Ada-ECE (%) \u2193",
        "AUROC \u2191",
        "Brier Score \u2193",
    ]
    hdr_row = start_row + 1
    ws.row_dimensions[hdr_row].height = 22
    for col_idx, h in enumerate(headers, start=1):
        cell = get_cell(ws, hdr_row, col_idx)
        cell.value = h
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER if col_idx > 2 else ALIGN_LEFT
        cell.border = BORDER_HEADER
    return hdr_row


def format_dataset_table(
    ws: Worksheet,
    start_row: int,
    model_name: str,
    dataset_name: str,
    arch_data: dict[str, pd.DataFrame],
) -> int:
    """Writes a single model's results table for a dataset, returning the next available row."""
    hdr_row = write_table_headers(ws, start_row, f"Model: {model_name}", f"Dataset: {dataset_name}")
    bench_df = arch_data["bench"]
    ump_df = arch_data["umpire"]
    b_sub: Any = bench_df[bench_df["dataset"] == dataset_name]
    u_sub: Any = ump_df[ump_df["dataset"] == dataset_name]

    cur_row = hdr_row + 1
    for idx, (disp_m, raw_m, src, regime) in enumerate(TARGET_METHODS):
        fill = FILL_ZEBRA if idx % 2 == 1 else FILL_WHITE
        ws.row_dimensions[cur_row].height = 19

        c1 = get_cell(ws, cur_row, 1)
        c2 = get_cell(ws, cur_row, 2)
        c3 = get_cell(ws, cur_row, 3)
        c4 = get_cell(ws, cur_row, 4)
        c5 = get_cell(ws, cur_row, 5)
        c6 = get_cell(ws, cur_row, 6)

        c1.value = disp_m
        c2.value = regime
        for c in [c1, c2, c3, c4, c5, c6]:
            c.font = FONT_DATA
            c.fill = fill
            c.border = BORDER_THIN

        c1.alignment = ALIGN_LEFT
        c2.alignment = ALIGN_LEFT

        if src == "bench":
            row_data: Any = b_sub[b_sub["method"] == raw_m]
            if len(row_data) > 0:
                r = row_data.iloc[0]
                c3.value = float(r["ece"])
                c3.number_format = "0.00%"
                c3.alignment = ALIGN_RIGHT

                c4.value = float(r["adaptive_ece"])
                c4.number_format = "0.00%"
                c4.alignment = ALIGN_RIGHT

                c5.value = float(r["auroc"])
                c5.number_format = "0.000"
                c5.alignment = ALIGN_RIGHT

                c6.value = float(r["brier"])
                c6.number_format = "0.0000"
                c6.alignment = ALIGN_RIGHT
        else:
            row_data = u_sub[u_sub["method"] == raw_m]
            if len(row_data) > 0:
                r = row_data.iloc[0]
                c3.value = float(r["cece"])
                c3.number_format = "0.00%"
                c3.alignment = ALIGN_RIGHT

                c4.value = "-"
                c4.alignment = ALIGN_CENTER

                c5.value = float(r["auc"])
                c5.number_format = "0.000"
                c5.alignment = ALIGN_RIGHT

                c6.value = "-"
                c6.alignment = ALIGN_CENTER

        cur_row += 1

    return cur_row + 2


def format_macro_table(
    ws: Worksheet,
    start_row: int,
    model_name: str,
    datasets: list[str],
    m_offset: int,
) -> int:
    """Writes dynamic Excel formula macro averages across all 14 datasets for one model."""
    hdr_row = write_table_headers(
        ws,
        start_row,
        f"Macro-Average Results: {model_name}",
        f"Averaged across {len(datasets)} Benchmarks",
    )
    cur_row = hdr_row + 1
    for idx, (disp_m, _raw_m, src, regime) in enumerate(TARGET_METHODS):
        fill = FILL_ZEBRA if idx % 2 == 1 else FILL_WHITE
        ws.row_dimensions[cur_row].height = 19
        src_row = m_offset + idx

        c1 = get_cell(ws, cur_row, 1)
        c2 = get_cell(ws, cur_row, 2)
        c3 = get_cell(ws, cur_row, 3)
        c4 = get_cell(ws, cur_row, 4)
        c5 = get_cell(ws, cur_row, 5)
        c6 = get_cell(ws, cur_row, 6)

        c1.value = disp_m
        c2.value = regime
        for c in [c1, c2, c3, c4, c5, c6]:
            c.font = FONT_DATA
            c.fill = fill
            c.border = BORDER_THIN

        c1.alignment = ALIGN_LEFT
        c2.alignment = ALIGN_LEFT

        ece_refs = ", ".join([f"'{d}'!C{src_row}" for d in datasets])
        c3.value = f"=AVERAGE({ece_refs})"
        c3.number_format = "0.00%"
        c3.alignment = ALIGN_RIGHT

        auc_refs = ", ".join([f"'{d}'!E{src_row}" for d in datasets])
        c5.value = f"=AVERAGE({auc_refs})"
        c5.number_format = "0.000"
        c5.alignment = ALIGN_RIGHT

        if src == "bench":
            aece_refs = ", ".join([f"'{d}'!D{src_row}" for d in datasets])
            c4.value = f"=AVERAGE({aece_refs})"
            c4.number_format = "0.00%"
            c4.alignment = ALIGN_RIGHT

            brier_refs = ", ".join([f"'{d}'!F{src_row}" for d in datasets])
            c6.value = f"=AVERAGE({brier_refs})"
            c6.number_format = "0.0000"
            c6.alignment = ALIGN_RIGHT
        else:
            c4.value = "-"
            c4.alignment = ALIGN_CENTER
            c6.value = "-"
            c6.alignment = ALIGN_CENTER

        cur_row += 1

    return cur_row + 2


def adjust_column_widths(ws: Worksheet) -> None:
    """Sets professional column widths with safety padding."""
    col_widths = {1: 28, 2: 30, 3: 17, 4: 19, 5: 16, 6: 19}
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width


def build_workbook(
    data: dict[str, dict[str, pd.DataFrame]],
    output_path: Path,
    include_m3: bool = True,
    include_mqt: bool = True,
) -> None:
    """Builds complete Excel workbook with Macro_Average sheet and 14 dataset sheets."""
    wb = openpyxl.Workbook()
    datasets = sorted(data["m3"]["bench"]["dataset"].unique())

    ws_macro: Any = wb.active
    assert ws_macro is not None
    ws_macro.title = "Macro_Average"
    ws_macro.views.sheetView[0].showGridLines = True

    for d in datasets:
        ws_d = wb.create_sheet(title=d)
        ws_d.views.sheetView[0].showGridLines = True
        cur_r = 2
        if include_m3:
            cur_r = format_dataset_table(ws_d, cur_r, "M3-LLaVA (7B)", d, data["m3"])
        if include_mqt:
            _ = format_dataset_table(ws_d, cur_r, "MQT-LLaVA (7B)", d, data["mqt"])
        adjust_column_widths(ws_d)

    cur_macro_r = 2
    if include_m3:
        cur_macro_r = format_macro_table(
            ws_macro, cur_macro_r, "M3-LLaVA (7B)", datasets, m_offset=4
        )
    if include_mqt:
        mqt_offset = 19 if include_m3 else 4
        _ = format_macro_table(
            ws_macro, cur_macro_r, "MQT-LLaVA (7B)", datasets, m_offset=mqt_offset
        )
    adjust_column_widths(ws_macro)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"Successfully generated workbook: {output_path}")


def main() -> None:
    repo_root = Path(__file__).resolve().parent.parent
    data = load_all_data(repo_root)

    out_master = repo_root / "results/calibration_results_master.xlsx"
    build_workbook(data, out_master, include_m3=True, include_mqt=True)

    out_m3 = repo_root / "results/calibration_results_m3_llava.xlsx"
    build_workbook(data, out_m3, include_m3=True, include_mqt=False)

    out_mqt = repo_root / "results/calibration_results_mqt_llava.xlsx"
    build_workbook(data, out_mqt, include_m3=False, include_mqt=True)


if __name__ == "__main__":
    main()
