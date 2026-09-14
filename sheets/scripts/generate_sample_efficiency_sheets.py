#!/usr/bin/env python3
"""
Generate Sample Efficiency and Small-Dataset Benchmark Excel Workbook.

Produces sheets/sample_efficiency_results.xlsx containing:
1. Macro_Mean: Multi-budget (N in {50, 100, 200, 500, 1000, 1500, Full}) macro averages
   for M3-LLaVA and MQT-LLaVA across the 5 evaluated methods.
2. 14 Individual Dataset Sheets (ai2d, chartqa, ..., vqav2):
   Detailed scaling data by budget tier, method, and architecture.
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any, cast

import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
RESULTS_DIR = ROOT / "results" / "experiments" / "sample_efficiency"
SHEETS_DIR = ROOT / "sheets"
SHEETS_DIR.mkdir(parents=True, exist_ok=True)
DEFAULT_OUTPUT_FILE = SHEETS_DIR / "sample_efficiency_results.xlsx"

ALL_14_DATASETS: list[str] = [
    "ai2d",
    "chartqa",
    "docvqa",
    "gqa",
    "infographicvqa",
    "lego-puzzles",
    "mmbench",
    "mmmu",
    "pope",
    "scienceqa",
    "seedbench",
    "textvqa",
    "vizwiz-vqa",
    "vqav2",
]

DEFAULT_BUDGETS: list[str] = ["50", "100", "200", "500", "1000", "1500", "Full"]

METHODS_ORDER: list[str] = [
    "Naive Confidence (NC)",
    "Temperature Scaling (TS)",
    "Platt Scaling (1D)",
    "Trajectory Platt (5D)",
    "VCPS-5D (Our Method)",
]


def parse_args() -> argparse.Namespace:
    """CLI argument parser for sheet generation."""
    parser = argparse.ArgumentParser(description="Generate Sample Efficiency Excel Workbook.")
    parser.add_argument(
        "--results_dir",
        type=str,
        default=str(RESULTS_DIR),
        help="Path to sample efficiency experiment results directory.",
    )
    parser.add_argument(
        "--output_file",
        type=str,
        default=str(DEFAULT_OUTPUT_FILE),
        help="Path to output Excel workbook.",
    )
    return parser.parse_args()


def load_summaries(
    results_dir: Path,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Loads M3 and MQT summary and macro CSVs."""
    m3_summary_path = results_dir / "sample_efficiency_m3_summary.csv"
    mqt_summary_path = results_dir / "sample_efficiency_mqt_summary.csv"
    m3_macro_path = results_dir / "sample_efficiency_m3_macro.csv"
    mqt_macro_path = results_dir / "sample_efficiency_mqt_macro.csv"

    if not m3_summary_path.exists() or not mqt_summary_path.exists():
        raise FileNotFoundError(
            f"Missing summary CSVs in {results_dir.resolve()}. "
            "Please run scripts/run_sample_efficiency_study.py first."
        )

    df_m3 = pd.read_csv(m3_summary_path)
    df_mqt = pd.read_csv(mqt_summary_path)
    df_macro_m3 = pd.read_csv(m3_macro_path) if m3_macro_path.exists() else pd.DataFrame()
    df_macro_mqt = pd.read_csv(mqt_macro_path) if mqt_macro_path.exists() else pd.DataFrame()

    return df_m3, df_mqt, df_macro_m3, df_macro_mqt


def build_macro_mean_table(
    df_macro_m3: pd.DataFrame,
    df_macro_mqt: pd.DataFrame,
) -> pd.DataFrame:
    """Builds consolidated Macro_Mean sheet comparing M3 and MQT across budgets."""
    rows: list[dict[str, Any]] = []

    arch_configs = [
        ("M3-LLaVA (7B)", df_macro_m3),
        ("MQT-LLaVA (7B)", df_macro_mqt),
    ]

    for model_name, df_macro in arch_configs:
        if df_macro.empty:
            continue
        budgets_in_df = [
            b for b in DEFAULT_BUDGETS if str(b) in df_macro["budget"].astype(str).values
        ]
        for budget in budgets_in_df:
            sub = cast(pd.DataFrame, df_macro[df_macro["budget"].astype(str) == str(budget)])
            for m in METHODS_ORDER:
                sub_m = cast(pd.DataFrame, sub[sub["method"] == m])
                if sub_m.empty:
                    continue
                rec = sub_m.iloc[0]
                rows.append(
                    {
                        "Model": model_name,
                        "Budget Tier": str(budget),
                        "Method": m,
                        "Adaptive ECE (%)": round(float(rec["ada_ece_mean"]), 4),
                        "Ada-ECE Std (%)": round(float(rec["ada_ece_std"]), 4),
                        "ECE (%)": round(float(rec["ece_mean"]), 4),
                        "ECE Std (%)": round(float(rec["ece_std"]), 4),
                        "AUROC": round(float(rec["auroc_mean"]), 4),
                        "AUROC Std": round(float(rec["auroc_std"]), 4),
                        "Brier Score": round(float(rec["brier_mean"]), 4),
                        "Brier Std": round(float(rec["brier_std"]), 4),
                        "Mean Sample Size": round(float(rec["sample_size_mean"]), 1),
                    }
                )

    return pd.DataFrame(rows)


def build_dataset_table(
    df_m3: pd.DataFrame,
    df_mqt: pd.DataFrame,
    dataset_name: str,
) -> pd.DataFrame:
    """Builds scaling performance table for a specific benchmark dataset."""
    rows: list[dict[str, Any]] = []

    arch_configs = [
        ("M3-LLaVA (7B)", df_m3),
        ("MQT-LLaVA (7B)", df_mqt),
    ]

    for model_name, df_arch in arch_configs:
        df_ds = cast(pd.DataFrame, df_arch[df_arch["dataset"] == dataset_name])
        if df_ds.empty:
            continue
        budgets_in_df = [b for b in DEFAULT_BUDGETS if str(b) in df_ds["budget"].astype(str).values]
        for budget in budgets_in_df:
            sub = cast(pd.DataFrame, df_ds[df_ds["budget"].astype(str) == str(budget)])
            for m in METHODS_ORDER:
                sub_m = cast(pd.DataFrame, sub[sub["method"] == m])
                if sub_m.empty:
                    continue
                rec = sub_m.iloc[0]
                effective_n = round(float(rec["effective_n_mean"]))
                is_capped = bool(rec.get("is_capped", False))
                sample_display = f"{effective_n} (Full)" if is_capped else str(effective_n)

                rows.append(
                    {
                        "Model": model_name,
                        "Budget Tier": str(budget),
                        "Method": m,
                        "Sample Size": sample_display,
                        "Adaptive ECE (%)": round(float(rec["adaptive_ece_percent_mean"]), 4),
                        "Ada-ECE Std (%)": round(float(rec["adaptive_ece_percent_std"]), 4),
                        "ECE (%)": round(float(rec["ece_percent_mean"]), 4),
                        "ECE Std (%)": round(float(rec["ece_percent_std"]), 4),
                        "AUROC": round(float(rec["auroc_mean"]), 4),
                        "AUROC Std": round(float(rec["auroc_std"]), 4),
                        "Brier Score": round(float(rec["brier_mean"]), 4),
                        "Brier Std": round(float(rec["brier_std"]), 4),
                    }
                )

    return pd.DataFrame(rows)


def write_macro_sheet_with_notes(
    writer: pd.ExcelWriter,
    sheet_name: str,
    table_df: pd.DataFrame,
) -> None:
    """Writes Macro_Mean table and explanatory documentation to Excel sheet."""
    table_df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=0)
    notes = [
        "Macro_Mean Calibration Scaling Notes:",
        "1. Metrics represent the unweighted macro-average across all 14 vision-language benchmarks.",
        "2. Standard deviations reflect stability across 5 independent stratified subsampling seeds (42, 43, 44, 45, 46).",
        "3. For budgets exceeding a benchmark's training split, the sample size was capped at len(train_df) to prevent dataset-dropout confounding.",
        "4. Directionality: Ada-ECE, ECE, and Brier (lower is better); AUROC (higher is better).",
    ]
    notes_df = pd.DataFrame({"Notes": notes})
    notes_df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=len(table_df) + 2)


def write_dataset_sheet_with_notes(
    writer: pd.ExcelWriter,
    sheet_name: str,
    table_df: pd.DataFrame,
    dataset_name: str,
) -> None:
    """Writes per-dataset scaling table and notes to Excel sheet."""
    table_df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=0)
    notes = [
        f"Benchmark: {dataset_name.upper()}",
        "1. Mean and Std are computed across 5 random subsampling seeds for budgets N < Full.",
        "2. 'Sample Size' indicates actual training instances used during post-hoc calibrator fitting.",
        "3. Values tagged '(Full)' indicate the training partition was smaller than the budget tier and capped.",
    ]
    notes_df = pd.DataFrame({"Notes": notes})
    notes_df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=len(table_df) + 2)


def apply_excel_styling(writer: pd.ExcelWriter) -> None:
    """Applies professional formatting: auto column widths, header fill, and pane freeze."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(name="Calibri", size=11, bold=True, color="000000")
    header_fill = PatternFill(start_color="EAECEE", end_color="EAECEE", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    book = getattr(writer, "book", None)
    if book is None:
        return

    for ws in book.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                val = str(cell.value) if cell.value is not None else ""
                if len(val) > 50:
                    continue
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 13)


def generate_sheets(
    results_dir: Path | str = RESULTS_DIR,
    output_file: Path | str = DEFAULT_OUTPUT_FILE,
) -> Path:
    """Loads summaries, builds tables, and writes all 15 sheets to Excel."""
    res_dir = Path(results_dir)
    out_file = Path(output_file)
    out_file.parent.mkdir(parents=True, exist_ok=True)

    df_m3, df_mqt, df_macro_m3, df_macro_mqt = load_summaries(res_dir)
    macro_table = build_macro_mean_table(df_macro_m3, df_macro_mqt)

    try:
        with pd.ExcelWriter(out_file, engine="openpyxl") as writer:
            write_macro_sheet_with_notes(writer, "Macro_Mean", macro_table)
            for ds in ALL_14_DATASETS:
                ds_table = build_dataset_table(df_m3, df_mqt, ds)
                write_dataset_sheet_with_notes(writer, ds, ds_table, ds)
            apply_excel_styling(writer)
        return out_file
    except PermissionError:
        fallback_file = out_file.with_name(f"{out_file.stem}_updated.xlsx")
        with pd.ExcelWriter(fallback_file, engine="openpyxl") as writer:
            write_macro_sheet_with_notes(writer, "Macro_Mean", macro_table)
            for ds in ALL_14_DATASETS:
                ds_table = build_dataset_table(df_m3, df_mqt, ds)
                write_dataset_sheet_with_notes(writer, ds, ds_table, ds)
            apply_excel_styling(writer)
        return fallback_file


def main() -> None:
    """Main execution function to build Excel workbook."""
    args = parse_args()
    print(f"Generating Excel workbook from {args.results_dir}...")
    out_path = generate_sheets(args.results_dir, args.output_file)
    print(f"Successfully generated {out_path} with 15 sheets.")


if __name__ == "__main__":
    main()
